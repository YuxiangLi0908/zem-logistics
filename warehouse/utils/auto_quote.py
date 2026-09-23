"""Address import, immutable batch creation and durable queue operations."""
import hashlib
import math
import re
import uuid
from datetime import date, timedelta
from decimal import Decimal, InvalidOperation

from django.db import transaction
from django.db.models import Count
from django.utils import timezone
from openpyxl import load_workbook

from warehouse.models.auto_quote import AutoQuoteAddress, AutoQuoteBatch, AutoQuoteItem, AutoQuoteWorkerState
from warehouse.utils.multi_carrier_quote import prepare_quote
from warehouse.utils.quote_analysis_storage import get_profile, index_item

GROUPS = ("LA", "SAV", "NJ")
TERMINAL = ("success", "partial", "failed", "no_quote", "cancelled")
RETRYABLE = ("partial", "failed", "no_quote", "cancelled")


def check_group(value):
    if value not in GROUPS:
        raise ValueError("请选择 LA、SAV 或 NJ 地址组")
    return value


def clean_text(value):
    return " ".join(str(value or "").strip().split())


def read_addresses(upload):
    if not upload.name.lower().endswith(".xlsx") or upload.size > 5 * 1024 * 1024:
        raise ValueError("请上传不超过5MB的 .xlsx 文件")
    workbook = load_workbook(upload, read_only=True, data_only=False)
    try:
        sheet = workbook.active
        rows = sheet.iter_rows(min_col=2, max_col=6, values_only=True)
        headers = tuple(clean_text(v) for v in next(rows, ()))
        if headers[:4] != ("城市", "州", "邮编", "详细地址") or len(headers) != 5 or headers[4].replace(" ", "").replace("（", "(").replace("）", ")").lower() != "距仓库距离(miles)":
            raise ValueError("B～F列表头应为：城市、州、邮编、详细地址、距仓库距离（miles）")
        addresses, errors = [], []
        for row_number, row in enumerate(rows, 2):
            if row_number > 10001:
                raise ValueError("单次最多导入10000行地址")
            if all(v is None or str(v).strip() == "" for v in row):
                continue
            try:
                if any(isinstance(v, str) and v.startswith("=") for v in row):
                    raise ValueError("请将公式转换为实际值")
                city, state, zipcode, address, distance = row
                city, state, address = clean_text(city), clean_text(state).upper(), clean_text(address)
                if isinstance(zipcode, (int, float)) and not isinstance(zipcode, bool):
                    if not math.isfinite(zipcode) or int(zipcode) != zipcode:
                        raise ValueError("邮编格式错误")
                    zipcode = str(int(zipcode)).zfill(5)
                zipcode = clean_text(zipcode)
                if not city or len(city) > 120 or not address or len(address) > 500:
                    raise ValueError("城市和详细地址必填，长度分别不能超过120、500字符")
                if not re.fullmatch(r"[A-Z]{2}", state):
                    raise ValueError("州必须是两位字母")
                if not re.fullmatch(r"\d{5}(-\d{4})?", zipcode):
                    raise ValueError("邮编应为5位数字或 ZIP+4")
                miles = None if distance is None or str(distance).strip() == "" else Decimal(str(distance))
                if miles is not None:
                    if not miles.is_finite() or miles < 0 or miles >= Decimal("100000000"):
                        raise ValueError("距离必须是有效的非负数字")
                    miles = miles.quantize(Decimal("0.01"))
                    if miles >= Decimal("100000000"):
                        raise ValueError("距离数值过大")
                fingerprint = hashlib.sha256("|".join([city.casefold(), state, zipcode, address.casefold()]).encode()).hexdigest()
                addresses.append(dict(city=city, state=state, zipcode=zipcode, address=address,
                                      distance_miles=miles, fingerprint=fingerprint))
            except (ValueError, InvalidOperation, TypeError) as exc:
                errors.append(f"第{row_number}行：{exc}")
        return addresses, errors
    finally:
        workbook.close()


@transaction.atomic
def import_addresses(group, rows, user_id):
    check_group(group)
    import_id, added = uuid.uuid4(), 0
    for row in rows:
        _, created = AutoQuoteAddress.objects.get_or_create(
            group=group, fingerprint=row["fingerprint"],
            defaults={**row, "uploaded_by_id": user_id, "import_id": import_id},
        )
        added += int(created)
    return {"added": added, "duplicates": len(rows) - added, "import_id": str(import_id)}


def address_payload(parameters, address):
    return {**parameters, "destinationWarehouse": f'{address["city"]}, {address["state"]} {address["zipcode"]}',
            "destinationCity": address["city"], "destinationState": address["state"],
            "destinationPostCode": address["zipcode"], "destinationDetailAddress": address["address"]}


def validate_parameters(parameters, sample):
    prepare_quote(address_payload(parameters, sample))
    if date.fromisoformat(parameters["pickupDate"]) < timezone.localdate():
        raise ValueError("取件日期不能早于今天")
    for key in ("originWarehouse", "originCity", "originState", "originPostCode"):
        if not isinstance(parameters.get(key), str) or not parameters[key].strip():
            raise ValueError("请完整选择发货仓库")
    if int(parameters.get("quoteType", 1)) not in (1, 2):
        raise ValueError("运输类型不正确")
    if int(parameters.get("destinationType", 1)) not in (1, 2, 3):
        raise ValueError("收货地址类型不正确")
    declared = float(parameters["declaredValue"])
    if not math.isfinite(declared) or declared <= 0:
        raise ValueError("申报价值必须大于0")
    for item in parameters["items"]:
        for key in ("length", "width", "height", "weight", "pieces", "palletCount"):
            value = float(item[key])
            if not math.isfinite(value) or value <= 0 or (key in ("pieces", "palletCount") and not value.is_integer()):
                raise ValueError("货物尺寸、重量必须大于0，件数和板数必须是正整数")


def queue_lock():
    AutoQuoteWorkerState.objects.get_or_create(pk=1)
    return AutoQuoteWorkerState.objects.select_for_update().get(pk=1)


@transaction.atomic
def create_batch(group, parameters, user_id, submission_id, parent=None):
    check_group(group)
    token = uuid.UUID(str(submission_id))
    queue_lock()
    existing = AutoQuoteBatch.objects.filter(submission_id=token, operator_id=user_id).first()
    if existing:
        return existing
    if parent:
        if parent.items.filter(status__in=("pending", "running")).exists():
            raise ValueError("请等待原任务结束后再重试")
        snapshots = list(parent.items.filter(status__in=RETRYABLE).values_list("address_snapshot", flat=True))
    else:
        snapshots = [row.snapshot() for row in AutoQuoteAddress.objects.filter(group=group)]
    if not snapshots:
        raise ValueError("没有可询价的地址")
    validate_parameters(parameters, snapshots[0])
    # A second tab cannot launch the same conditions while they are still running.
    for active in AutoQuoteBatch.objects.filter(group=group, operator_id=user_id, status__in=("queued", "running")):
        if active.parameters == parameters:
            return active
    batch = AutoQuoteBatch.objects.create(group=group, parameters=parameters, profile=get_profile(parameters), operator_id=user_id,
                                         submission_id=token, parent=parent)
    AutoQuoteItem.objects.bulk_create([AutoQuoteItem(batch=batch, address_snapshot=a) for a in snapshots])
    return batch


def finish_batch(batch_id):
    batch = AutoQuoteBatch.objects.get(pk=batch_id)
    if not batch.items.filter(status__in=("pending", "running")).exists():
        AutoQuoteBatch.objects.filter(pk=batch_id).update(
            status="stopped" if batch.stop_requested else "completed", finished_at=timezone.now())


@transaction.atomic
def stop_batch(batch):
    queue_lock()
    AutoQuoteBatch.objects.filter(pk=batch.pk).update(stop_requested=True)
    batch.items.filter(status="pending").update(status="cancelled", finished_at=timezone.now())
    finish_batch(batch.pk)


@transaction.atomic
def claim_item(concurrency=3):
    state = queue_lock()
    now = timezone.now()
    state.heartbeat_at = now
    state.save(update_fields=["heartbeat_at"])
    # A worker has at most 240s to execute; a 300s lease permits restart recovery.
    stale = list(AutoQuoteItem.objects.filter(status="running", lease_until__lt=now).values_list("pk", "batch_id", "batch__stop_requested"))
    for pk, batch_id, stopped in stale:
        AutoQuoteItem.objects.filter(pk=pk).update(status="cancelled" if stopped else "pending", lease_token=None, lease_until=None)
        if stopped:
            finish_batch(batch_id)
    if AutoQuoteItem.objects.filter(status="running", lease_until__gte=now).count() >= concurrency:
        return None
    item = AutoQuoteItem.objects.filter(status="pending", batch__stop_requested=False).select_related("batch").first()
    if not item:
        return None
    item.status = "running"
    item.lease_token = uuid.uuid4()
    item.lease_until = now + timedelta(seconds=300)
    item.started_at = item.started_at or now
    item.attempts += 1
    item.save(update_fields=["status", "lease_token", "lease_until", "started_at", "attempts"])
    AutoQuoteBatch.objects.filter(pk=item.batch_id).update(status="running")
    return item


def save_checkpoint(item, result, quote_uuid, payload):
    count = AutoQuoteItem.objects.filter(pk=item.pk, lease_token=item.lease_token, status="running").update(
        result=result, quote_uuid=quote_uuid, request_payload=payload)
    if not count:
        raise RuntimeError("询价任务已由其他执行进程接管")


def price_rows(payload):
    if isinstance(payload, dict):
        for key in ("rates", "quotes"):
            if isinstance(payload.get(key), list):
                return [row for row in payload[key] if isinstance(row, dict)]
        for value in payload.values():
            rows = price_rows(value)
            if rows:
                return rows
    return []


def classify_result(result):
    carriers = result.get("results", {})
    active = [carriers.get(name, {}) for name in ("maersk", "kakas")]
    def has_price(carrier):
        for row in price_rows(carrier):
            value = row.get("TotalQuote", row.get("totalPrice", row.get("price")))
            try:
                if value is not None and value != "" and math.isfinite(float(value)) and float(value) >= 0:
                    return True
            except (ValueError, TypeError):
                pass
        return False

    has_quotes = [has_price(carrier) for carrier in active]
    issues = [carrier.get("warning") or carrier.get("error") or
              ("平台询价失败" if carrier.get("status") != "success" else "") for carrier in active]
    if any(has_quotes):
        status = "success" if all(has_quotes) and not any(issues) else "partial"
    else:
        status = "failed" if any(issues) else "no_quote"
    return status, "；".join(str(issue) for issue in issues if issue)


@transaction.atomic
def complete_item(item, status, error=""):
    queue_lock()
    updated = AutoQuoteItem.objects.filter(pk=item.pk, lease_token=item.lease_token, status="running").update(
        status=status, error=error, finished_at=timezone.now(), lease_until=None)
    if updated:
        index_item(item.pk)
    finish_batch(item.batch_id)


def batch_summary(batch):
    counts = dict(batch.items.values("status").annotate(n=Count("id")).values_list("status", "n"))
    return {"id": batch.pk, "group": batch.group, "status": batch.status, "counts": counts,
            "total": sum(counts.values()), "stop_requested": batch.stop_requested,
            "origin": batch.parameters.get("originWarehouse", ""), "created_at": batch.created_at,
            "profile_code": f"AQ{batch.profile_id:06d}" if batch.profile_id else "待归档",
            "profile_id": batch.profile_id,
            "operator": (batch.operator.get_full_name() or batch.operator.get_username()) if batch.operator else "原用户已删除 / 未记录",
            "finished_at": batch.finished_at, "parent_id": batch.parent_id}
