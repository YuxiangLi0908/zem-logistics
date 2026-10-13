import csv
import json
import logging
from uuid import uuid4
from io import BytesIO
from datetime import date, timedelta
from zipfile import BadZipFile

from django.core.exceptions import PermissionDenied
from django.core.paginator import Paginator
from django.db import IntegrityError
from django.http import HttpResponse, JsonResponse
from django.shortcuts import get_object_or_404, render
from django.utils import timezone
from openpyxl.utils.exceptions import InvalidFileException
from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from urllib.parse import quote

from warehouse.models.auto_quote import AutoQuoteAddress, AutoQuoteBatch, AutoQuoteWorkerState
from warehouse.models.system_parameter import SystemParameter
from warehouse.utils.auto_quote import (
    batch_summary, check_group, create_batch, import_addresses, read_addresses, stop_batch,
)
from warehouse.utils.multi_carrier_quote import gateway_config
from warehouse.utils.quote_analysis import analysis_options, build_analysis
from warehouse.utils.quote_destination_history import destination_history
from warehouse.utils.auto_quote_lifecycle import schedule_worker

logger = logging.getLogger(__name__)

def visible_batches(user):
    queryset = AutoQuoteBatch.objects.select_related("operator")
    return queryset if user.is_staff else queryset.filter(operator=user)


def auto_quote_get(request):
    if not request.user.is_authenticated:
        raise PermissionDenied
    if request.GET.get("step") == "auto_quote_history":
        return render(request, "post_port/new_sop/leader_check/auto_quote_history.html")
    if request.GET.get("step") == "auto_quote_analysis":
        return render(request, "post_port/new_sop/leader_check/auto_quote_analysis.html")
    try:
        kind = request.GET.get("kind", "batches")
        if kind == "destination_history":
            return JsonResponse({"success": True, **destination_history(visible_batches(request.user), request.GET)})
        if kind == "analysis_options":
            return JsonResponse({"success": True, **analysis_options(visible_batches(request.user))})
        if kind == "analysis":
            return JsonResponse({"success": True, **build_analysis(visible_batches(request.user), request.GET)})
        if kind == "analysis_export":
            report = build_analysis(visible_batches(request.user), request.GET, export=True)
            response = HttpResponse(content_type="text/csv; charset=utf-8")
            response["Content-Disposition"] = f'attachment; filename="price-analysis-{report["profile"]["code"]}.csv"'
            response.write("\ufeff")
            writer = csv.writer(response)
            writer.writerow(["比较编号", "取件开始日期", "取件结束日期", "币种", "包含部分报价", "地址", "平台", "承运商", "服务", "最新取件日期", "最新价格",
                             "上次取件日期", "上次价格", "涨跌金额", "涨跌%", "期间涨跌%", "最低", "最高", "均价", "CV%", "振幅%", "最大相邻涨跌%", "有效日", "采样日", "覆盖率%"])
            def safe_cell(value):
                if value is None:
                    return ""
                if isinstance(value, str) and value.startswith(("=", "+", "-", "@", "\t", "\r")):
                    return "'" + value
                return value
            for row in report["rows"]:
                values = [report["profile"]["code"], report["filters"]["start"], report["filters"]["end"],
                          report["filters"]["currency"], report["filters"]["include_partial"]]
                values += [row[key] for key in ("address", "platform", "carrier", "service", "latest_date", "latest", "previous_date", "previous", "latest_change", "latest_change_pct",
                                               "period_change_pct", "minimum", "maximum", "mean", "volatility_pct", "range_pct", "max_adjacent_move_pct", "samples", "expected", "coverage_pct")]
                writer.writerow([safe_cell(value) for value in values])
            return response
        if request.GET.get("step") == "auto_quote_export":
            batch = get_object_or_404(visible_batches(request.user), pk=request.GET.get("batch"))
            response = HttpResponse(content_type="text/csv; charset=utf-8")
            response["Content-Disposition"] = f'attachment; filename="auto-quote-{batch.pk}.csv"'
            response.write("\ufeff")
            writer = csv.writer(response)
            writer.writerow(["组", "发货仓", "取件日期", "城市", "州", "邮编", "详细地址", "距离(miles)", "状态", "询价时间", "完成时间", "错误", "报价结果(JSON)"])

            def safe(value):
                text = str(value if value is not None else "")
                return "'" + text if text.startswith(("=", "+", "-", "@", "\t", "\r")) else text

            for item in batch.items.iterator():
                a = item.address_snapshot
                values = [batch.group, batch.parameters.get("originWarehouse"), batch.parameters.get("pickupDate"),
                          a["city"], a["state"], a["zipcode"], a["address"], a.get("distance_miles"),
                          item.status, item.started_at, item.finished_at, item.error,
                          json.dumps(item.result, ensure_ascii=False)]
                writer.writerow([safe(value) for value in values])
            return response
        if kind == "addresses":
            group = check_group(request.GET.get("group"))
            page = Paginator(AutoQuoteAddress.objects.filter(group=group), 50).get_page(request.GET.get("page"))
            return JsonResponse({"success": True, "rows": [a.snapshot() for a in page],
                                 "page": page.number, "pages": page.paginator.num_pages, "total": page.paginator.count})
        if kind == "batch_copy":
            batch = get_object_or_404(visible_batches(request.user), pk=request.GET.get("batch"))
            # One read for the full result table; omit large request payloads not used in the clipboard.
            records = list(batch.items.values("id", "address_snapshot", "status", "started_at", "finished_at", "error", "result"))
            for row in records:
                row["address"] = row.pop("address_snapshot")
            return JsonResponse({"success": True, "batch_id": batch.pk, "rows": records})
        if kind == "batch":
            batch = get_object_or_404(visible_batches(request.user), pk=request.GET.get("batch"))
            items = batch.items.all()
            if request.GET.get("status"):
                items = items.filter(status=request.GET["status"])
            page = Paginator(items, 30).get_page(request.GET.get("page"))
            return JsonResponse({"success": True, "batch": batch_summary(batch), "parameters": batch.parameters,
                                 "page": page.number, "pages": page.paginator.num_pages,
                                 "rows": [{"id": item.pk, "address": item.address_snapshot, "status": item.status,
                                           "started_at": item.started_at, "finished_at": item.finished_at,
                                           "error": item.error, "result": item.result, "request_payload": item.request_payload,
                                           "attempts": item.attempts} for item in page]})
        batches = visible_batches(request.user)
        for key, lookup in (("start", "gte"), ("end", "lte")):
            if request.GET.get(key):
                batches = batches.filter(**{f"parameters__pickupDate__{lookup}": date.fromisoformat(request.GET[key]).isoformat()})
        if request.GET.get("start") and request.GET.get("end") and request.GET["start"] > request.GET["end"]:
            raise ValueError("开始日期不能晚于结束日期")
        page = Paginator(batches, 20).get_page(request.GET.get("page"))
        state = AutoQuoteWorkerState.objects.filter(pk=1).first()
        online = bool(state and state.heartbeat_at and state.heartbeat_at > timezone.now() - timedelta(seconds=30))
        starting = bool(state and state.run_token and state.lease_until and state.lease_until > timezone.now() and not online)
        pending = visible_batches(request.user).filter(items__status__in=("pending", "running")).exists()
        return JsonResponse({"success": True, "batches": [batch_summary(batch) for batch in page],
                             "page": page.number, "pages": page.paginator.num_pages, "worker_online": online,
                             "worker_starting": starting, "worker_error": state.last_error if state else "",
                             "worker_needed": pending,
                             "groups": {group: AutoQuoteAddress.objects.filter(group=group).count() for group in ("LA", "SAV", "NJ")}})
    except (ValueError, TypeError) as exc:
        return JsonResponse({"success": False, "message": str(exc)}, status=400)


def auto_quote_post(request):
    if not request.user.is_authenticated:
        raise PermissionDenied
    step = None
    try:
        step = request.POST.get("step")
        if step == "auto_quote_table_export":
            title = request.POST.get("title")
            if title not in ("距仓库距离分析", "收货地址价格分析"):
                raise ValueError("请选择有效的分析表")
            upload = request.FILES.get("sheets_file")
            try:
                sheets = json.loads(upload.read().decode("utf-8") if upload else request.POST.get("sheets") or "[]")
            except (UnicodeDecodeError, json.JSONDecodeError) as exc:
                raise ValueError("导出数据不是有效的 UTF-8 JSON，请刷新页面后重新导出") from exc
            if not isinstance(sheets, list) or not 1 <= len(sheets) <= 2:
                raise ValueError("导出表格格式错误")
            workbook = Workbook(write_only=True)
            alignment = Alignment(vertical="top", wrap_text=True)
            heading_font = Font(bold=True)
            for index, rows in enumerate(sheets):
                if not isinstance(rows, list) or not rows:
                    raise ValueError("导出表格内容不能为空")
                sheet = workbook.create_sheet(title if index == 0 else "承运商明细")
                widths = {}
                for row in rows:
                    if not isinstance(row, list) or len(row) > 30 or any(not isinstance(cell, str) for cell in row):
                        raise ValueError("导出表格格式错误")
                    for column, value in enumerate(row, 1):
                        widths[column] = max(widths.get(column, 0), max(map(len, value.split("\n"))))
                sheet.freeze_panes = "A2"
                if widths:
                    sheet.auto_filter.ref = f"A1:{get_column_letter(max(widths))}{len(rows)}"
                for column, width in widths.items():
                    sheet.column_dimensions[get_column_letter(column)].width = min(60, max(16, width + 2))
                for row_number, row in enumerate(rows, 1):
                    cells = []
                    for value in row:
                        cell = WriteOnlyCell(sheet, value=ILLEGAL_CHARACTERS_RE.sub("", value))
                        # All exported cells are display text, including addresses and postal codes.
                        cell.data_type = "s"
                        cell.alignment = alignment
                        if row_number == 1:
                            cell.font = heading_font
                        cells.append(cell)
                    sheet.append(cells)
            output = BytesIO()
            workbook.save(output)
            response = HttpResponse(output.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
            response["Content-Disposition"] = "attachment; filename*=UTF-8''" + quote(title + ".xlsx")
            return response
        if step == "auto_quote_wake":
            if visible_batches(request.user).filter(items__status__in=("pending", "running")).exists():
                schedule_worker()
            return JsonResponse({"success": True})
        if step == "auto_quote_import":
            group = check_group(request.POST.get("group"))
            upload = request.FILES.get("file")
            if not upload:
                raise ValueError("请选择Excel文件")
            rows, errors = read_addresses(upload)
            if errors:
                return JsonResponse({"success": False, "message": f"发现{len(errors)}行错误，本次未导入，请修正后重传。",
                                     "errors": errors[:100]}, status=400)
            if not rows:
                raise ValueError("Excel中没有地址记录")
            return JsonResponse({"success": True, **import_addresses(group, rows, request.user.pk)})
        if step == "auto_quote_start":
            parameters = json.loads(request.POST.get("quote_payload") or "{}")
            if not isinstance(parameters, dict):
                raise ValueError("询价参数格式错误")
            # Resolve the selected origin on the server, exactly as for the selector.
            origin = next((a for a in SystemParameter.get_zem_warehouse_addresses()
                           if a["warehouse"] == parameters.get("originWarehouse")), None)
            if not origin:
                raise ValueError("请选择有效的发货仓库")
            parameters.update(originCity=origin.get("city"), originState=origin.get("state"),
                              originPostCode=origin.get("postCode"), originDetailAddress=origin.get("detailAddress", ""),
                              originType=1)
            for key in ("destinationWarehouse", "destinationCity", "destinationState", "destinationPostCode", "destinationDetailAddress"):
                parameters.pop(key, None)
            gateway_config()
            batch = create_batch(request.POST.get("group"), parameters, request.user.pk, request.POST.get("submission_id"))
            schedule_worker()
            return JsonResponse({"success": True, "batch": batch_summary(batch)})
        batch = get_object_or_404(visible_batches(request.user), pk=request.POST.get("batch"))
        if step == "auto_quote_stop":
            stop_batch(batch)
            batch.refresh_from_db()
        elif step == "auto_quote_retry":
            gateway_config()
            batch = create_batch(batch.group, batch.parameters, request.user.pk, request.POST.get("submission_id"), parent=batch)
            schedule_worker()
        else:
            raise ValueError("未知操作")
        return JsonResponse({"success": True, "batch": batch_summary(batch)})
    except (ValueError, TypeError, KeyError, OverflowError, AttributeError, BadZipFile, InvalidFileException) as exc:
        return JsonResponse({"success": False, "message": str(exc)}, status=400)
    except IntegrityError:
        return JsonResponse({"success": False, "message": "重复提交，请刷新任务列表查看"}, status=409)
    except Exception as exc:
        if step != "auto_quote_table_export":
            raise
        error_id = uuid4().hex[:12]
        logger.exception("Analysis Excel export failed, error_id=%s", error_id)
        reason = "服务器生成 Excel 时发生异常"
        if isinstance(exc, OSError):
            reason = f"服务器文件读写失败：{exc.strerror or '请检查临时目录权限及磁盘空间'}"
        elif isinstance(exc, MemoryError):
            reason = "服务器内存不足，请缩小筛选范围后重试"
        return JsonResponse({"success": False, "message": f"{reason}（{type(exc).__name__}，错误编号：{error_id}）"}, status=500)
